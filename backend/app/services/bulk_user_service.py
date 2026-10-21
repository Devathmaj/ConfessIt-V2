"""Import participant accounts from CSV/XLSX with all-or-nothing validation.

Input size, row count, workbook contents, and fields are validated before writes; account
capacity is checked in the insert transaction. Formula cells, external links, macros,
encrypted entries, and malformed CSV are refused rather than evaluated or repaired.
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass, replace
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Iterable, Iterator, Sequence

from fastapi import HTTPException, UploadFile, status
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import func, or_, select, text
from sqlalchemy.exc import IntegrityError

from .. import limits as app_limits
from ..accounts import AccountType, account_cap, account_cap_lock_key
from ..db_models import UserRecord
from ..logger import get_logger
from ..models import AdminUserCreate, _reject_unchangeable_characters
from ..postgres import Database

logger = get_logger(__name__)

# Bound on rows iterated from the sheet/file, including blank rows. A file with 1000 data
# rows plus a handful of separators is ordinary; thousands of blank rows before the data is
# not, and iterating them is work an operator can still be tricked into paying for with an
# enormous file. The row limit above is the product rule; this one is the parser's own safety
# bound and is deliberately more generous.
_MAX_PHYSICAL_ROWS = app_limits.BULK_UPLOAD_MAX_ROWS * 2 + 100

# The XLSX container is a ZIP, so its advertised sizes are attacker-controlled and the
# parser must inspect them before decompression. A normal workbook is a few hundred KB
# uncompressed; 25 MB is far past any legitimate 1000-row file.
_XLSX_MAX_ENTRIES = 2000
_XLSX_MAX_UNCOMPRESSED_BYTES = 25 * 1024 * 1024
_XLSX_MAX_COMPRESSION_RATIO = 200
_XLSX_MAX_COLUMNS = 20

# Reasons an XLSX part is refused even though it is a valid ZIP member. None of these are
# needed to hold cell text, and each is a known place to hide a capability the upload has no
# business carrying: VBA macros, external workbook references, embedded OLE/ActiveX objects
# and custom ribbon UI.
_FORBIDDEN_XLSX_PREFIXES = (
    "xl/externallinks/",
    "xl/embeddings/",
    "xl/activex/",
    "xl/oleobjects/",
    "xl/macrosheets/",
    "xl/querytables/",
    "xl/connections",
    "xl/vbaproject",
    "customui/",
)

# The editable identity fields, in the order the UI presents them. The first row of a file
# may use any of the aliases below and may put them in any order; these canonical names are
# what the parser reports back.
_FIELD_ORDER: tuple[str, ...] = (
    "registration_number",
    "email",
    "full_name",
    "gender",
    "year",
    "department",
)

_FIELD_LABELS: dict[str, str] = {
    "registration_number": "registration number",
    "email": "email",
    "full_name": "full name",
    "gender": "gender",
    "year": "year",
    "department": "department",
}

_FIELD_ALIASES: dict[str, frozenset[str]] = {
    "registration_number": frozenset(
        {
            "registration_number",
            "registration_no",
            "registrationnumber",
            "registrationnum",
            "registration",
            "register_number",
            "register_no",
            "reg_number",
            "regno",
            "reg_no",
        }
    ),
    "email": frozenset(
        {"email", "email_address", "emailaddress", "email_id", "mail"}
    ),
    "full_name": frozenset({"full_name", "fullname", "name", "student_name"}),
    "gender": frozenset({"gender", "sex"}),
    "year": frozenset(
        {"year", "year_of_study", "study_year", "current_year"}
    ),
    "department": frozenset(
        {"department", "department_division", "division", "dept"}
    ),
}

_ALIAS_TO_FIELD: dict[str, str] = {
    alias: field for field, aliases in _FIELD_ALIASES.items() for alias in aliases
}

# A value starting with one of these is treated as a spreadsheet formula rather than as
# text. Identity fields have no legitimate reason to begin with any of them, and refusing the
# value at import is cheaper than remembering to neutralise it at every future export.
_FORMULA_PREFIXES = ("=", "+", "-", "@")

# How many individual problems are quoted back. A 1000-row file can be wrong in 1000 ways,
# and a response that lists all of them is a response that is harder to read than the file.
_MAX_REPORTED_PROBLEMS = 25
_MAX_PROBLEM_LENGTH = 240

_READ_CHUNK_BYTES = 64 * 1024

_YEAR_TEXT_RE = re.compile(r"^[0-9]{1,3}$")
_MULTIPLE_SEPARATORS_RE = re.compile(r"_+")


class BulkUploadError(HTTPException):
    """A refusal that names every problem it found, not just the first one.

    ``detail`` is a plain string rather than an object because that is the shape every
    existing client error path already renders (``describeApiError``). Individual row
    problems are appended to it with their sheet row number and field label, so the operator
    can fix the file in one pass instead of resubmitting once per mistake.
    """

    def __init__(
        self,
        status_code: int,
        message: str,
        problems: Sequence["RowProblem"] = (),
    ) -> None:
        detail = message
        if problems:
            rendered = [problem.render() for problem in problems[:_MAX_REPORTED_PROBLEMS]]
            detail = f"{message} {'; '.join(rendered)}"
            remaining = len(problems) - len(rendered)
            if remaining > 0:
                detail += f"; and {remaining} more problem(s)."
        super().__init__(status_code=status_code, detail=detail)


@dataclass(frozen=True)
class RowProblem:
    """One reason a row (or the heading row) was refused."""

    row: int
    field: str
    message: str

    def render(self) -> str:
        message = " ".join(self.message.split())[:_MAX_PROBLEM_LENGTH]
        return f"row {self.row} {self.field}: {message}"


@dataclass
class ParticipantUploadRow:
    """A fully validated participant row, ready to be written."""

    row: int
    regno: str
    name: str
    email: str
    gender: str
    year: int
    department: str


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _safe_label(value: Any) -> str:
    """A heading rendered back in an error without its control characters."""
    text = " ".join(str(value or "").split())
    return text[:60] or "(blank)"


def _problem(
    problems: list[RowProblem], row: int, field: str, message: str
) -> None:
    problems.append(
        RowProblem(
            row=row,
            field=_FIELD_LABELS.get(field, field),
            message=message,
        )
    )


def _normalise_heading(value: Any) -> str:
    """Reduce a heading to the spelling the alias table uses.

    This is intentionally forgiving about presentation (case, spaces, hyphens, a trailing
    asterisk from a "required" marker) and unforgiving about meaning: an unrecognised
    heading is refused, not guessed at.
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return ""
    if not isinstance(value, str):
        value = str(value)
    text = value.strip().lower().lstrip("\ufeff")
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return _MULTIPLE_SEPARATORS_RE.sub("_", text).strip("_")


def _prepare_headings(values: Sequence[Any], row: int) -> list[str]:
    """Map the heading row to canonical field names, or refuse the file."""
    problems: list[RowProblem] = []
    seen: dict[str, int] = {}
    fields: list[str] = []

    for index, raw in enumerate(values, start=1):
        normalised = _normalise_heading(raw)
        if not normalised:
            _problem(problems, row, f"column {index}", "heading is blank")
            continue
        field = _ALIAS_TO_FIELD.get(normalised)
        if field is None:
            _problem(
                problems,
                row,
                f"column {index}",
                f"unknown column {_safe_label(raw)!r}; participants have only "
                + ", ".join(_FIELD_LABELS[name] for name in _FIELD_ORDER),
            )
            continue
        if field in seen:
            _problem(
                problems,
                row,
                _FIELD_LABELS[field],
                f"appears in both column {seen[field]} and column {index}",
            )
            continue
        seen[field] = index
        fields.append(field)

    for field in _FIELD_ORDER:
        if field not in seen:
            _problem(problems, row, _FIELD_LABELS[field], "required column is missing")

    if problems:
        raise BulkUploadError(
            status.HTTP_400_BAD_REQUEST,
            "The heading row is not valid, so nothing was uploaded.",
            problems,
        )
    return fields


def _text_value(
    value: Any,
    field: str,
    row: int,
    problems: list[RowProblem],
    *,
    allow_numeric: bool = False,
) -> str | None:
    """Turn one cell into stripped text, or record why it cannot be text.

    ``allow_numeric`` is only for the registration number: Excel stores what looks like a
    number as a number, and refusing ``12345`` there would make a perfectly ordinary register
    number impossible to upload. Every other text field requires an actual string, because
    silently stringifying a date or a boolean produces a plausible-looking but wrong account.
    """
    if _is_blank(value):
        return None

    if isinstance(value, bool):
        _problem(problems, row, field, "must be text, not true/false")
        return None

    if isinstance(value, (int, float)):
        if allow_numeric and not (isinstance(value, float) and not value.is_integer()):
            return str(int(value))
        _problem(problems, row, field, "must be text")
        return None

    if isinstance(value, (date, datetime, time)):
        _problem(problems, row, field, "must be text, not a date or time")
        return None

    if not isinstance(value, str):
        _problem(problems, row, field, "must be text")
        return None

    text = value.strip()
    if not text:
        return None

    try:
        _reject_unchangeable_characters(text, _FIELD_LABELS[field])
    except ValueError as exc:
        _problem(problems, row, field, str(exc))
        return None

    # The shared character rule permits tabs and line breaks because confessions are prose
    # and may legitimately contain them. An identity field is one line: a newline in a name
    # is either a broken file or the first half of a CSV-injection payload, and neither is a
    # value the account columns are meant to hold.
    if any(character in "\t\r\n\u2028\u2029" for character in text):
        _problem(problems, row, field, "must be a single line of text")
        return None

    if text.startswith(_FORMULA_PREFIXES):
        _problem(
            problems,
            row,
            field,
            "must not start with =, +, - or @ (spreadsheet formula prefixes are refused)",
        )
        return None

    return text


def _year_value(value: Any, row: int, problems: list[RowProblem]) -> int | None:
    """Read a year as a whole number.

    The range itself is enforced by ``AdminUserCreate`` (and by the CHECK constraint under
    it), so this only answers the earlier question of whether the cell is a number at all.
    ``"02"`` normalises to ``2`` rather than becoming a second cohort spelling.
    """
    if _is_blank(value):
        return None
    if isinstance(value, bool):
        _problem(problems, row, "year", "must be a whole number")
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if value.is_integer():
            return int(value)
        _problem(problems, row, "year", "must be a whole number")
        return None
    if isinstance(value, str):
        text = value.strip()
        if not _YEAR_TEXT_RE.fullmatch(text):
            _problem(problems, row, "year", "must be a whole number")
            return None
        return int(text)
    _problem(problems, row, "year", "must be a whole number")
    return None


def _validate_row(
    row: int, raw: dict[str, Any], problems: list[RowProblem]
) -> ParticipantUploadRow | None:
    """Validate one data row against the same payload the single-create route uses."""
    registration_number = _text_value(
        raw.get("registration_number"),
        "registration_number",
        row,
        problems,
        allow_numeric=True,
    )
    full_name = _text_value(raw.get("full_name"), "full_name", row, problems)
    email = _text_value(raw.get("email"), "email", row, problems)
    gender = _text_value(raw.get("gender"), "gender", row, problems)
    department = _text_value(raw.get("department"), "department", row, problems)
    year = _year_value(raw.get("year"), row, problems)

    # Required columns are named individually. An empty cell is a missing answer; a cell that
    # failed conversion has already been reported above and must not also be called "required".
    for field, value, raw_value in (
        ("registration_number", registration_number, raw.get("registration_number")),
        ("full_name", full_name, raw.get("full_name")),
        ("email", email, raw.get("email")),
        ("gender", gender, raw.get("gender")),
        ("year", year, raw.get("year")),
        ("department", department, raw.get("department")),
    ):
        if _is_blank(raw_value):
            _problem(problems, row, field, "is required")

    if registration_number is not None and len(registration_number) > app_limits.USER_REGNO_MAX:
        _problem(
            problems,
            row,
            "registration_number",
            f"must be at most {app_limits.USER_REGNO_MAX} characters",
        )
    if full_name is not None and len(full_name) > app_limits.USER_NAME_MAX:
        _problem(
            problems,
            row,
            "full_name",
            f"must be at most {app_limits.USER_NAME_MAX} characters",
        )
    if department is not None and len(department) > app_limits.USER_DEPARTMENT_MAX:
        _problem(
            problems,
            row,
            "department",
            f"must be at most {app_limits.USER_DEPARTMENT_MAX} characters",
        )
    # EmailStr checks syntax but not the width of the column underneath it.
    if email is not None and len(email) > 320:
        _problem(problems, row, "email", "must be at most 320 characters")

    if gender is not None:
        gender = gender.strip().lower()
    if email is not None:
        # The single-create route lowercases the address before storing it; do the same here
        # so the same person cannot exist twice under different casing.
        email = email.strip().lower()

    if any(
        value is None
        for value in (
            registration_number,
            full_name,
            email,
            gender,
            year,
            department,
        )
    ):
        return None

    try:
        validated = AdminUserCreate(
            account_type=AccountType.USERS,
            regno=registration_number,
            name=full_name,
            email=email,
            gender=gender,
            year=year,
            department=department,
        )
    except ValidationError as exc:
        for error in exc.errors():
            location = error.get("loc") or ()
            field = str(location[-1]) if location else "row"
            if field not in _FIELD_LABELS:
                field = "row"
            message = str(error.get("msg") or "is invalid").replace("Value error, ", "")
            _problem(problems, row, field, message)
        return None

    return ParticipantUploadRow(
        row=row,
        regno=validated.regno,
        name=validated.name,
        email=str(validated.email).strip().lower(),
        gender=validated.gender,
        year=validated.year,
        department=validated.department,
    )


def _parse_grid(
    rows: Iterable[tuple[int, Sequence[Any]]],
) -> list[ParticipantUploadRow]:
    """Read the heading and data rows, enforcing the row ceiling along the way."""
    problems: list[RowProblem] = []
    parsed: list[ParticipantUploadRow] = []
    heading_fields: list[str] | None = None
    heading_row = 0
    data_rows = 0
    physical_rows = 0

    for row_number, cells in rows:
        physical_rows += 1
        if physical_rows > _MAX_PHYSICAL_ROWS:
            raise BulkUploadError(
                status.HTTP_400_BAD_REQUEST,
                f"The file has more than {_MAX_PHYSICAL_ROWS} rows (including blank rows), "
                "so it was not read further. Remove the empty rows and try again.",
            )

        values = list(cells)
        # Formatting can make a spreadsheet report thousands of empty trailing columns.
        # They are not data, so trailing empties are dropped before the row is interpreted.
        while values and _is_blank(values[-1]):
            values.pop()
        if not values or all(_is_blank(value) for value in values):
            continue

        if heading_fields is None:
            heading_row = row_number
            heading_fields = _prepare_headings(values, row_number)
            continue

        data_rows += 1
        if data_rows > app_limits.BULK_UPLOAD_MAX_ROWS:
            raise BulkUploadError(
                status.HTTP_400_BAD_REQUEST,
                f"At most {app_limits.BULK_UPLOAD_MAX_ROWS} participant rows may be uploaded "
                f"in one file; this file contains more than that. Split it into smaller "
                f"files. Nothing was uploaded.",
            )

        if len(values) > len(heading_fields):
            _problem(
                problems,
                row_number,
                "columns",
                f"has {len(values)} values but the heading row has "
                f"{len(heading_fields)} columns",
            )
            continue

        raw = {
            field: values[index] if index < len(values) else None
            for index, field in enumerate(heading_fields)
        }
        participant = _validate_row(row_number, raw, problems)
        if participant is not None:
            parsed.append(participant)

    if heading_fields is None:
        raise BulkUploadError(
            status.HTTP_400_BAD_REQUEST,
            f"The file is empty: row {heading_row or 1} must contain the column headings.",
        )
    if data_rows == 0:
        raise BulkUploadError(
            status.HTTP_400_BAD_REQUEST,
            "The file has a heading row but no participant rows, so nothing was uploaded.",
        )
    if problems:
        raise BulkUploadError(
            status.HTTP_400_BAD_REQUEST,
            "The file was not uploaded because some rows are invalid.",
            problems,
        )

    _check_file_duplicates(parsed)
    return parsed


def _check_file_duplicates(rows: Sequence[ParticipantUploadRow]) -> None:
    """Refuse duplicate registration numbers or addresses inside one file.

    Registration numbers are compared case-insensitively for this check: the database's
    uniqueness rule is case-sensitive, but ``23CS1`` and ``23cs1`` are a duplicate to every
    person reading the file, and the second one would be an account nobody intended to make.
    """
    problems: list[RowProblem] = []
    seen_regno: dict[str, ParticipantUploadRow] = {}
    seen_email: dict[str, ParticipantUploadRow] = {}

    for row in rows:
        regno_key = row.regno.casefold()
        previous = seen_regno.get(regno_key)
        if previous is not None:
            _problem(
                problems,
                row.row,
                "registration_number",
                f"duplicates row {previous.row}",
            )
        else:
            seen_regno[regno_key] = row

        email_key = row.email.lower()
        previous = seen_email.get(email_key)
        if previous is not None:
            _problem(problems, row.row, "email", f"duplicates row {previous.row}")
        else:
            seen_email[email_key] = row

    if problems:
        raise BulkUploadError(
            status.HTTP_400_BAD_REQUEST,
            "The file was not uploaded because it contains duplicate participants.",
            problems,
        )


def _detect_csv_delimiter(text: str) -> str:
    """Choose comma, semicolon or tab from the first non-empty line.

    Excel in a comma-decimal locale writes semicolon-separated CSV, and it would be unhelpful
    to call that malformed. Counting occurrences on the heading line is enough for a
    six-column file and is deterministic, unlike a general-purpose sniffer that can be fooled
    by punctuation inside a name.
    """
    first_line = next((line for line in text.splitlines() if line.strip()), "")
    counts = {delimiter: first_line.count(delimiter) for delimiter in (",", ";", "\t")}
    delimiter = max(counts, key=counts.get)
    return delimiter if counts[delimiter] else ","


def _read_csv(content: bytes) -> Iterator[tuple[int, list[Any]]]:
    """Decode a CSV file and yield ``(row number, cells)`` without trusting its extension."""
    if b"\x00" in content:
        raise BulkUploadError(
            status.HTTP_400_BAD_REQUEST,
            "The CSV file contains a NUL byte, so it is not a text file. Nothing was uploaded.",
        )
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise BulkUploadError(
            status.HTTP_400_BAD_REQUEST,
            "The CSV file is not valid UTF-8 text. Save it as UTF-8 CSV and try again.",
        ) from None
    if not text.strip():
        raise BulkUploadError(
            status.HTTP_400_BAD_REQUEST,
            "The CSV file is empty. Nothing was uploaded.",
        )

    reader = csv.reader(
        io.StringIO(text, newline=""),
        delimiter=_detect_csv_delimiter(text),
        strict=True,
    )
    try:
        for row_number, cells in enumerate(reader, start=1):
            yield row_number, cells
    except csv.Error as exc:
        raise BulkUploadError(
            status.HTTP_400_BAD_REQUEST,
            f"The CSV file is malformed near row {reader.line_num}: {exc}. Nothing was uploaded.",
        ) from None


def _inspect_xlsx_archive(content: bytes) -> None:
    """Validate the ZIP container before openpyxl is allowed to parse it.

    ``xlsx`` is a ZIP of XML parts. The checks here are the ones a generic spreadsheet parser
    does not make for you: entry count and total expanded size (a decompression bomb), a sane
    compression ratio, no encrypted members, no path traversal in member names, and no macro
    or external-workbook parts pretending to be ordinary content.
    """
    if not content.startswith(b"PK\x03\x04"):
        raise BulkUploadError(
            status.HTTP_400_BAD_REQUEST,
            "The file is named .xlsx but is not an XLSX workbook. Nothing was uploaded.",
        )

    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
            if not infos:
                raise BulkUploadError(
                    status.HTTP_400_BAD_REQUEST,
                    "The XLSX workbook is empty. Nothing was uploaded.",
                )
            if len(infos) > _XLSX_MAX_ENTRIES:
                raise BulkUploadError(
                    status.HTTP_400_BAD_REQUEST,
                    "The XLSX workbook contains too many parts to be a normal spreadsheet. "
                    "Nothing was uploaded.",
                )

            names: set[str] = set()
            total_uncompressed = 0
            for info in infos:
                name = info.filename.replace("\\", "/")
                lower = name.lower()

                if (
                    name.startswith("/")
                    or re.match(r"^[A-Za-z]:", name)
                    or any(part == ".." for part in name.split("/"))
                ):
                    raise BulkUploadError(
                        status.HTTP_400_BAD_REQUEST,
                        "The XLSX workbook contains an unsafe member path. Nothing was uploaded.",
                    )
                if lower in names:
                    raise BulkUploadError(
                        status.HTTP_400_BAD_REQUEST,
                        "The XLSX workbook contains duplicate members. Nothing was uploaded.",
                    )
                names.add(lower)

                if info.flag_bits & 0x1:
                    raise BulkUploadError(
                        status.HTTP_400_BAD_REQUEST,
                        "The XLSX workbook contains an encrypted part. Nothing was uploaded.",
                    )
                if info.compress_type not in (zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED):
                    raise BulkUploadError(
                        status.HTTP_400_BAD_REQUEST,
                        "The XLSX workbook uses an unsupported compression method. "
                        "Nothing was uploaded.",
                    )
                if any(lower.startswith(prefix) for prefix in _FORBIDDEN_XLSX_PREFIXES):
                    raise BulkUploadError(
                        status.HTTP_400_BAD_REQUEST,
                        "The XLSX workbook contains macros, external links or embedded "
                        "objects. Nothing was uploaded.",
                    )

                total_uncompressed += info.file_size
                if total_uncompressed > _XLSX_MAX_UNCOMPRESSED_BYTES:
                    raise BulkUploadError(
                        status.HTTP_400_BAD_REQUEST,
                        "The XLSX workbook expands to more data than a spreadsheet of this "
                        "size needs. Nothing was uploaded.",
                    )
                # A small compressed member that claims to expand enormously is the shape of
                # a bomb. Legitimate XML parts are nowhere near this ratio.
                if (
                    info.file_size > 2 * 1024 * 1024
                    and info.compress_size > 0
                    and info.file_size / info.compress_size > _XLSX_MAX_COMPRESSION_RATIO
                ):
                    raise BulkUploadError(
                        status.HTTP_400_BAD_REQUEST,
                        "The XLSX workbook has an unsafe compression ratio. Nothing was uploaded.",
                    )

            if "[content_types].xml" not in names or "xl/workbook.xml" not in names:
                raise BulkUploadError(
                    status.HTTP_400_BAD_REQUEST,
                    "The file is not a complete XLSX workbook. Nothing was uploaded.",
                )
    except (zipfile.BadZipFile, OSError, RuntimeError, NotImplementedError):
        raise BulkUploadError(
            status.HTTP_400_BAD_REQUEST,
            "The file is named .xlsx but is not a readable XLSX workbook. Nothing was uploaded.",
        ) from None


def _read_xlsx(content: bytes) -> Iterator[tuple[int, list[Any]]]:
    """Yield ``(row number, cells)`` from the first worksheet without evaluating formulas."""
    _inspect_xlsx_archive(content)

    try:
        workbook = load_workbook(
            io.BytesIO(content),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except Exception as exc:  # openpyxl/zip/xml parsers all surface different exceptions
        logger.warning(
            "Rejected an unreadable bulk-upload workbook (%s).", type(exc).__name__
        )
        raise BulkUploadError(
            status.HTTP_400_BAD_REQUEST,
            "The file is named .xlsx but could not be read as a valid workbook. "
            "Nothing was uploaded.",
        ) from None

    try:
        worksheets = workbook.worksheets
        if not worksheets:
            raise BulkUploadError(
                status.HTTP_400_BAD_REQUEST,
                "The XLSX workbook has no worksheet. Nothing was uploaded.",
            )
        worksheet = worksheets[0]
        # A malicious sheet can declare 16,384 columns; refusing an absurd width before
        # iterating is cheaper than trimming them row by row.
        max_column = int(worksheet.max_column or 1)
        if max_column > _XLSX_MAX_COLUMNS:
            raise BulkUploadError(
                status.HTTP_400_BAD_REQUEST,
                f"The XLSX workbook reports {max_column} columns, which is more than the "
                f"{_XLSX_MAX_COLUMNS} a participant sheet can need. Nothing was uploaded.",
            )

        try:
            for row_number, cells in enumerate(
                worksheet.iter_rows(
                    min_row=1,
                    max_row=_MAX_PHYSICAL_ROWS,
                    max_col=max_column,
                    values_only=False,
                ),
                start=1,
            ):
                values: list[Any] = []
                for cell in cells:
                    if cell.data_type == "f":
                        raise BulkUploadError(
                            status.HTTP_400_BAD_REQUEST,
                            f"Row {row_number} contains a formula. Supply the calculated "
                            "value as plain text instead. Nothing was uploaded.",
                        )
                    if cell.data_type == "e":
                        raise BulkUploadError(
                            status.HTTP_400_BAD_REQUEST,
                            f"Row {row_number} contains a spreadsheet error value. "
                            "Nothing was uploaded.",
                        )
                    values.append(cell.value)
                yield row_number, values
        except BulkUploadError:
            raise
        except Exception as exc:
            logger.warning(
                "Rejected an unreadable bulk-upload worksheet (%s).", type(exc).__name__
            )
            raise BulkUploadError(
                status.HTTP_400_BAD_REQUEST,
                "The XLSX worksheet could not be read. Nothing was uploaded.",
            ) from None
    finally:
        workbook.close()


async def read_bulk_upload(file: UploadFile) -> list[ParticipantUploadRow]:
    """Read and validate an upload, without touching the database.

    Reading is bounded by ``BULK_UPLOAD_MAX_BYTES`` regardless of the declared size, because
    ``UploadFile.size`` is supplied by the multipart parser and must not be the only limit.
    """
    filename = Path(str(file.filename or "")).name
    suffix = Path(filename).suffix.lower()
    if suffix not in (".csv", ".xlsx"):
        raise BulkUploadError(
            status.HTTP_400_BAD_REQUEST,
            "Only .csv and .xlsx files are accepted. Nothing was uploaded.",
        )

    max_bytes = int(app_limits.BULK_UPLOAD_MAX_BYTES)
    declared_size = getattr(file, "size", None)
    if declared_size is not None and declared_size > max_bytes:
        raise BulkUploadError(
            status.HTTP_413_CONTENT_TOO_LARGE,
            f"The file is larger than the {max_bytes // (1024 * 1024)} MB limit. "
            "Nothing was uploaded.",
        )

    chunks: list[bytes] = []
    total = 0
    while True:
        chunk = await file.read(_READ_CHUNK_BYTES)
        if not chunk:
            break
        total += len(chunk)
        if total > max_bytes:
            raise BulkUploadError(
                status.HTTP_413_CONTENT_TOO_LARGE,
                f"The file is larger than the {max_bytes // (1024 * 1024)} MB limit. "
                "Nothing was uploaded.",
            )
        chunks.append(chunk)

    content = b"".join(chunks)
    if not content:
        raise BulkUploadError(
            status.HTTP_400_BAD_REQUEST,
            "The uploaded file is empty. Nothing was uploaded.",
        )

    if suffix == ".csv":
        if content.startswith(b"PK\x03\x04"):
            raise BulkUploadError(
                status.HTTP_400_BAD_REQUEST,
                "The file is named .csv but contains a ZIP/XLSX workbook. Rename it to .xlsx "
                "or save it as CSV. Nothing was uploaded.",
            )
        rows = _read_csv(content)
    else:
        rows = _read_xlsx(content)

    return _parse_grid(rows)


def _reconcile_departments(
    session, rows: list[ParticipantUploadRow]
) -> list[ParticipantUploadRow]:
    """Fold uploaded departments onto existing spellings and enforce MAX_DEPARTMENTS.

    Departments are free text with no table of their own, so the ceiling is on distinct
    case-insensitive values. Joining an existing department costs nothing; only a name no row
    carries opens a slot. The whole upload is refused if the new names would cross the
    ceiling, rather than writing the rows that fit and dropping the rest.
    """
    existing_values = session.scalars(select(UserRecord.department).distinct()).all()
    existing_by_norm: dict[str, str] = {}
    for value in existing_values:
        if isinstance(value, str) and value.strip():
            existing_by_norm.setdefault(value.strip().lower(), value.strip())

    existing_norms = set(existing_by_norm)
    # Seeded with the existing spellings and then extended as the file is read, so
    # "CSE" and "cse" in one upload become one department with the first spelling seen --
    # exactly as they would if they had been created one at a time through an edit that
    # joined the existing value. Without this the upload could store two spellings of one
    # department while the distinct count treated them as one.
    planned_by_norm: dict[str, str] = dict(existing_by_norm)
    new_norms: set[str] = set()
    reconciled: list[ParticipantUploadRow] = []
    for row in rows:
        norm = row.department.strip().lower()
        planned = planned_by_norm.get(norm)
        if planned is not None:
            reconciled.append(
                replace(row, department=planned) if planned != row.department else row
            )
        else:
            planned_by_norm[norm] = row.department
            new_norms.add(norm)
            reconciled.append(row)

    cap = int(app_limits.MAX_DEPARTMENTS)
    if len(existing_norms) + len(new_norms) > cap:
        raise BulkUploadError(
            status.HTTP_409_CONFLICT,
            f"The upload was cancelled before any row was written because it would exceed "
            f"the department/division limit: {len(existing_norms)} already exist, the file "
            f"would add {len(new_norms)} new one(s), and the limit is {cap} "
            f"(MAX_DEPARTMENTS). Nothing was uploaded.",
        )
    return reconciled


def _reject_existing_conflicts(
    session, rows: Sequence[ParticipantUploadRow]
) -> None:
    """Refuse rows whose registration number or email is already in ``users``.

    Both columns are compared case-insensitively. The database's unique indexes are
    case-sensitive, but a registration number or address that differs only in case is a
    duplicate to every person reading the file, and the second row would be an account nobody
    intended to create. This is deliberately stricter than the single-create form, which only
    enforces the column's own rule; an upload is a bulk act and gets the stricter reading.
    """
    regnos = sorted({row.regno.lower() for row in rows})
    emails = sorted({row.email.lower() for row in rows})
    conflicts = session.execute(
        select(UserRecord.regno, UserRecord.email).where(
            or_(
                func.lower(UserRecord.regno).in_(regnos),
                func.lower(UserRecord.email).in_(emails),
            )
        )
    ).all()

    existing_regnos = {str(regno).lower() for regno, _ in conflicts}
    existing_emails = {str(email).lower() for _, email in conflicts}
    problems: list[RowProblem] = []
    for row in rows:
        if row.regno.lower() in existing_regnos:
            _problem(problems, row.row, "registration_number", "already exists")
        if row.email.lower() in existing_emails:
            _problem(problems, row.row, "email", "already exists")

    if problems:
        raise BulkUploadError(
            status.HTTP_409_CONFLICT,
            "The upload was cancelled before any row was written because some participants "
            "already exist.",
            problems,
        )


def create_participants_bulk(
    db: Database, rows: list[ParticipantUploadRow]
) -> dict[str, Any]:
    """Write every validated row in one transaction, or write none of them.

    The account count is read and the rows inserted while holding the same per-table advisory
    lock the single-create route uses. Without it, two simultaneous uploads (or an upload
    beside a create form) could each read "there is room" and together cross
    ``MAX_USER_COUNT`` -- which is exactly the limit the check exists to enforce.
    """
    if not rows:
        raise BulkUploadError(
            status.HTTP_400_BAD_REQUEST,
            "There are no participant rows to upload. Nothing was uploaded.",
        )

    account_type = AccountType.USERS.value
    cap = account_cap(account_type)

    try:
        with db.sessions.begin() as session:
            session.execute(
                text("SELECT pg_advisory_xact_lock(:lock_key)"),
                {"lock_key": account_cap_lock_key(account_type)},
            )

            existing_count = (
                session.scalar(select(func.count()).select_from(UserRecord)) or 0
            )
            projected_count = existing_count + len(rows)
            if projected_count > cap:
                raise BulkUploadError(
                    status.HTTP_409_CONFLICT,
                    f"The upload was cancelled before any row was written because it would "
                    f"exceed the participant account limit: {existing_count} participants "
                    f"already exist, the file contains {len(rows)}, which would make "
                    f"{projected_count}, and MAX_USER_COUNT is {cap}. Nothing was uploaded.",
                )

            rows = _reconcile_departments(session, rows)
            _reject_existing_conflicts(session, rows)

            for row in rows:
                session.add(
                    UserRecord(
                        regno=row.regno,
                        name=row.name,
                        email=row.email,
                        # A new participant has no picture, like one created through the
                        # single-account form. Both halves are named so the row cannot end up
                        # with half an identity.
                        profile_picture_public_id=None,
                        profile_picture_format=None,
                        gender=row.gender,
                        year=row.year,
                        department=row.department,
                    )
                )
            session.flush()
    except IntegrityError:
        # The pre-checks above should have caught every duplicate. This is the backstop for a
        # writer that did not take the advisory lock -- the transaction is rolled back, so no
        # row from the file survives, and the caller gets the same all-or-nothing refusal.
        raise BulkUploadError(
            status.HTTP_409_CONFLICT,
            "The upload was cancelled because a registration number or email was written "
            "by another request while it was being checked. Nothing was uploaded; try again.",
        ) from None

    # The admin list caches its totals for a few seconds. A successful upload is a write the
    # operator is about to look at, so the next list read must not answer with the count from
    # before it. Imported here rather than at module scope to keep the upload service
    # independent of the admin listing module until a write actually succeeds. Failures are
    # logged rather than raised: the rows are already committed, and a cache that could not
    # be cleared must not turn a successful upload into a response that says it failed.
    try:
        from .admin_service import invalidate_admin_user_count_cache

        invalidate_admin_user_count_cache()
    except Exception as exc:  # pragma: no cover - defensive, import is local and stable
        logger.warning("Could not clear the admin user count cache: %s", exc)

    logger.info(
        "Bulk participant upload applied: %s row(s), %s/%s participants.",
        len(rows),
        existing_count + len(rows),
        cap,
    )
    return {
        "created": len(rows),
        "total_users": existing_count + len(rows),
        "max_user_count": cap,
        "remaining_slots": max(0, cap - existing_count - len(rows)),
    }
